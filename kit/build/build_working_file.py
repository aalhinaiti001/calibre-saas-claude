"""Builds kit/Calibre_Verdict_Working_File.xlsx in the Calibre brand.

Implements product/METHOD.md: a locked rubric of 4 to 9 criteria, blind panel scores (2 to 5 panelists,
2 to 6 finalists), the divergence view, calibration with written reasons, the weighted result and the
90/180/365-day outcome register. Every derived number is a live formula.

Usage: python3 kit/build/build_working_file.py [out.xlsx]
Then recalculate (LibreOffice) so cached values exist, e.g. the xlsx skill's scripts/recalc.py.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE.parent / "Calibre_Verdict_Working_File.xlsx"
RUBRIC = json.loads((HERE / "default-finance-rubric.json").read_text())["criteria"]

# Calibre tokens, brand/HANDBOOK.html §03. No Daftar cream, rust, camel or monospace.
FOREST50, FOREST100, FOREST600, FOREST700, FOREST900 = "F0F5F3", "D5E3DE", "265147", "1C3D35", "0E211B"
STONE600, STONE500, WHITE = "394740", "596B66", "FFFFFF"
SERIF, SANS = "Lora", "Plus Jakarta Sans"
USE_STATEMENT = ("Structured hiring advisory for finance roles. Not a psychometric assessment. Does not predict "
                 "performance. The decision, and its consequences, rest with the employer.")

N_FIN, N_CRIT, N_PAN = 6, 9, 5          # METHOD.md §1 maxima
FIRST = 7                                # first data row on the K x C sheets
LAST = FIRST + N_FIN * N_CRIT - 1        # 60
PANEL = [col(6 + p) for p in range(N_PAN)]            # F..J  scores
BASIS = [col(6 + N_PAN + p) for p in range(N_PAN)]    # K..O  evidence basis
NOTES = [col(6 + 2 * N_PAN + p) for p in range(N_PAN)]  # P..T  evidence notes

thin = Side(style="thin", color=FOREST100)
rule = Side(style="medium", color=FOREST600)
F_BODY = Font(name=SANS, size=10, color=STONE600)
F_BOLD = Font(name=SANS, size=10, bold=True, color=FOREST700)
F_HEAD = Font(name=SANS, size=9, bold=True, color=WHITE)
F_NOTE = Font(name=SANS, size=9, italic=True, color=STONE500)
F_TITLE = Font(name=SERIF, size=18, color=FOREST900)
F_EYEBROW = Font(name=SANS, size=8, bold=True, color=FOREST600)
FILL_HEAD = PatternFill("solid", fgColor=FOREST900)
FILL_INPUT = PatternFill("solid", fgColor=FOREST50)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

wb = Workbook()
wb.remove(wb.active)


def sheet(name, heading, intro, widths, landscape=False):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws["A1"], ws["A1"].font = "CALIBRE BY DAFTAR", F_EYEBROW
    ws["A2"], ws["A2"].font = heading, F_TITLE
    ws["A3"], ws["A3"].font, ws["A3"].alignment = intro, F_NOTE, Alignment(wrap_text=False, vertical="top")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[col(i)].width = w
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def head(ws, row, labels, start=1):
    for i, text in enumerate(labels):
        c = ws.cell(row, start + i, text)
        c.font, c.fill, c.alignment = F_HEAD, FILL_HEAD, CENTER if i else WRAP
        c.border = Border(bottom=rule)


def put(ws, ref, value, font=F_BODY, fill=None, align=WRAP, fmt=None, border=True):
    c = ws[ref]
    c.value, c.font, c.alignment = value, font, align
    if fill:
        c.fill = fill
    if border:
        c.border = Border(bottom=thin)
    if fmt:
        c.number_format = fmt
    return c


def validate(ws, rng, **kw):
    dv = DataValidation(allow_blank=True, showErrorMessage=True, **kw)
    ws.add_data_validation(dv)
    dv.add(rng)


def rows_kc():
    """(row, finalist index 1..6, criterion index 1..9) for the K x C sheets."""
    for i in range(1, N_FIN + 1):
        for j in range(1, N_CRIT + 1):
            yield FIRST + (i - 1) * N_CRIT + (j - 1), i, j


RUB = "'01 Rubric'"
SET = "'00 Setup'"
SCO = "'02 Scores'"
DIV = "'03 Divergence'"
CAL = "'04 Calibration'"
RUB_ROW = lambda j: 7 + j          # criteria on rows 8..16
PAN_ROW = lambda p: 7 + p          # panelists on rows 8..12
FIN_ROW = lambda i: 15 + i         # finalists on rows 16..21

# ------------------------------------------------------------------ Read me
ws = sheet("Read me", "Calibre Verdict · working file",
           "One workbook per engagement. Complete the sheets in order. Nothing is scored before sheet 01 is locked.", [16, 30, 90])
r = 5
head(ws, r, ["Sheet", "What it is for", "Rule"])
for name, purpose, rule_text in [
    ("00 Setup", "Panel, finalists and engagement dates", "Two to five panelists, one of them the hiring lead. Two to six finalists, K1 to K6, in the order received."),
    ("01 Rubric", "The criteria, weights and anchors", "Four to nine criteria, whole-number weights totalling 100. Locked before any finalist is scored; never changed after the first score."),
    ("02 Scores", "Every panelist's score and evidence note", "Enter only after every panelist has submitted, or the scoring deadline has passed. Scores are never edited afterwards."),
    ("03 Divergence", "The spread across panelists", "Calculated. A spread of 2 is a divergence, 3 or 4 a contradiction. Scores of 0 (no evidence) are excluded from the spread."),
    ("04 Calibration", "The resolution of every flagged cell", "Enter an agreed score and a written reason for each flagged cell. The original scores on sheet 02 stay as they are."),
    ("05 Results", "Calibrated scores and the weighted result", "Calculated. Entry order only: never sort, colour or label finalists. The weighted result is advisory."),
    ("06 Outcomes", "The 90, 180 and 365-day follow-up", "Four questions at each checkpoint, recorded by initials only."),
]:
    r += 1
    put(ws, f"A{r}", name, F_BOLD); put(ws, f"B{r}", purpose); put(ws, f"C{r}", rule_text)

r += 2
put(ws, f"A{r}", "Legend", F_BOLD, border=False)
r += 1
put(ws, f"A{r}", "", fill=FILL_INPUT); put(ws, f"B{r}", "Shaded cells are for entry. Every other value is a formula; do not overwrite it.")
r += 2
head(ws, r, ["Score", "Meaning", "When to use it"])
for score, meaning, when in [
    (5, "Strong evidence, exceeds the requirement", "Anchor 5 met, evidenced in the record."),
    (4, "Above the requirement", "Between anchors 3 and 5."),
    (3, "Meets the requirement", "Anchor 3 met, evidenced in the record."),
    (2, "Partially meets", "Between anchors 1 and 3."),
    (1, "Does not meet", "Anchor 1 describes the evidence."),
    (0, "No evidence supplied", "Not assessable on the evidence provided. Recorded as a gap, never as a low score, and excluded from the spread."),
]:
    r += 1
    put(ws, f"A{r}", score, F_BOLD, align=CENTER); put(ws, f"B{r}", meaning); put(ws, f"C{r}", when)
r += 2
head(ws, r, ["Evidence basis", "Definition", ""])
for basis, definition in [
    ("Documented", "Written evidence in the file: CV detail, work product, a test result, a reference supplied by the client."),
    ("Interview-stated", "The candidate said it in interview; not corroborated elsewhere."),
    ("Inferred", "Not stated; inferred from role titles, tenure or employer. Named as inference in the Verdict."),
    ("None", "No evidence either way. Score 0 and record the gap."),
]:
    r += 1
    put(ws, f"A{r}", basis, F_BOLD); ws.merge_cells(f"B{r}:C{r}"); put(ws, f"B{r}", definition)
r += 2
put(ws, f"A{r}", "Example entry", F_BOLD, border=False)
r += 1
ws.merge_cells(f"B{r}:C{r}")
put(ws, f"A{r}", "K1 · C2", F_BODY)
put(ws, f"B{r}", "Score 3 · Documented · \"CV and second interview note of 12 May: owned the monthly close for a five-entity group. Does not establish listed-group reporting.\" Evidence notes run to at least 40 characters, naming the source and the fact.")
r += 2
put(ws, f"A{r}", "Not in scope", F_BOLD, border=False)
for item in ["Candidate sourcing, interview scheduling, reference checks, offer negotiation",
             "Psychometric, personality, work-style, culture or fit assessment",
             "Automated candidate ranking, or any recommendation generated by a model",
             "The hiring decision itself"]:
    r += 1
    ws.merge_cells(f"B{r}:C{r}"); put(ws, f"B{r}", item, border=False)
r += 2
put(ws, f"A{r}", "Fonts", F_BOLD, border=False)
ws.merge_cells(f"B{r}:C{r}")
put(ws, f"B{r}", "Set in Lora and Plus Jakarta Sans (both free, SIL Open Font License). Install them, or Excel substitutes another face.", F_NOTE, border=False)
r += 2
put(ws, f"A{r}", "Use statement", F_BOLD, border=False)
ws.merge_cells(f"B{r}:C{r}")
put(ws, f"B{r}", USE_STATEMENT, F_BOLD, border=False)
for rr in range(6, r + 1):
    ws.row_dimensions[rr].height = None

# ------------------------------------------------------------------ 00 Setup
ws = sheet("00 Setup", "00 · Setup", "The panel, the finalists and the dates. Names are used on the working sheets; the Verdict uses initials.", [8, 34, 30, 16, 30])
put(ws, "A5", "Engagement", F_BOLD, border=False)
for rr, label in zip(range(6, 7), ["Engagement reference"]):
    put(ws, f"B{rr}", label, F_BOLD); put(ws, f"C{rr}", None, fill=FILL_INPUT)
head(ws, 7, ["Ref", "Panelist", "Title", "Hiring lead", "Email"])
for p in range(1, N_PAN + 1):
    rr = PAN_ROW(p)
    put(ws, f"A{rr}", f"P{p}", F_BOLD)
    for c in "BCDE":
        put(ws, f"{c}{rr}", None, fill=FILL_INPUT)
validate(ws, f"D{PAN_ROW(1)}:D{PAN_ROW(N_PAN)}", type="list", formula1='"Yes,No"')
put(ws, "B13", "Panelists named", F_BOLD); put(ws, "C13", f"=COUNTA(B{PAN_ROW(1)}:B{PAN_ROW(N_PAN)})")
put(ws, "D13", f'=IF(AND(C13>=2,C13<=5,COUNTIF(D{PAN_ROW(1)}:D{PAN_ROW(N_PAN)},"Yes")=1),"OK","Needs 2 to 5 panelists and exactly one hiring lead")', F_BOLD)
head(ws, 15, ["Ref", "Initials", "Name (purged 90 days after issue)", "Notice sent", "Source"])
for i in range(1, N_FIN + 1):
    rr = FIN_ROW(i)
    put(ws, f"A{rr}", f"K{i}", F_BOLD)
    for c in "BCE":
        put(ws, f"{c}{rr}", None, fill=FILL_INPUT)
    put(ws, f"D{rr}", None, fill=FILL_INPUT, fmt="d mmm yyyy")
put(ws, "B22", "Finalists named", F_BOLD); put(ws, "C22", f"=COUNTA(B{FIN_ROW(1)}:B{FIN_ROW(N_FIN)})")
put(ws, "D22", '=IF(AND(C22>=2,C22<=6),"OK","Needs 2 to 6 finalists")', F_BOLD)
head(ws, 24, ["", "Date", "", "", ""])
for rr, label in zip(range(25, 29), ["Rubric locked", "Scoring deadline", "Calibration session", "Verdict issued"]):
    put(ws, f"B{rr}", label, F_BOLD); put(ws, f"C{rr}", None, fill=FILL_INPUT, fmt="d mmm yyyy")
put(ws, "B29", "Personal data purge due", F_BOLD); put(ws, "C29", '=IF(C28="","",C28+90)', fmt="d mmm yyyy")

# ------------------------------------------------------------------ 01 Rubric
ws = sheet("01 Rubric", "01 · Rubric", "Amend wording, must or preferred, and weights at intake. Four to nine criteria. Leave unused rows blank. Once locked, nothing here changes.",
           [7, 32, 12, 9, 40, 40, 40], landscape=True)
put(ws, "A5", "Locked on", F_BOLD); put(ws, "B5", f'=IF({SET}!C25="","",{SET}!C25)', fmt="d mmm yyyy")
head(ws, 7, ["Ref", "Criterion", "Must / Preferred", "Weight %", "Anchor 1: does not meet", "Anchor 3: meets", "Anchor 5: strong evidence"])
for j, c in enumerate(RUBRIC, 1):
    rr = RUB_ROW(j)
    put(ws, f"A{rr}", c["ref"], F_BOLD)
    put(ws, f"B{rr}", c["name"], fill=FILL_INPUT)
    put(ws, f"C{rr}", c["must"], fill=FILL_INPUT)
    put(ws, f"D{rr}", c["weight"], fill=FILL_INPUT, align=CENTER)
    for k, key in zip("EFG", ["anchor_1", "anchor_3", "anchor_5"]):
        put(ws, f"{k}{rr}", c[key], fill=FILL_INPUT)
validate(ws, "C8:C16", type="list", formula1='"Must,Preferred"')
validate(ws, "D8:D16", type="whole", operator="between", formula1="1", formula2="100")
put(ws, "B17", "Total", F_BOLD); put(ws, "D17", "=SUM(D8:D16)", F_BOLD, align=CENTER)
put(ws, "B18", "Criteria in use", F_BOLD); put(ws, "D18", "=COUNTA(B8:B16)", align=CENTER)
put(ws, "B19", "Lock check", F_BOLD)
ws.merge_cells("C19:G19")
put(ws, "C19", '=IF(D18<4,"Not lockable: fewer than 4 criteria",IF(D18>9,"Not lockable: more than 9 criteria",'
               'IF(SUMPRODUCT((B8:B16<>"")*(D8:D16=""))>0,"Not lockable: a named criterion has no weight",'
               'IF(SUMPRODUCT((B8:B16="")*(D8:D16<>""))>0,"Not lockable: a weight has no criterion",'
               'IF(D17<>100,"Not lockable: weights total "&D17&", not 100","Ready to lock")))))', F_BOLD)
put(ws, "B21", "Source", F_NOTE, border=False)
ws.merge_cells("C21:G21")
put(ws, "C21", "Default finance rubric and anchors from the Calibre kit v1.0 working file. No criterion may touch personality, health or a protected characteristic.", F_NOTE, border=False)
ws.freeze_panes = "C8"


# ------------------------------------------------------------------ K x C helpers
def kc_labels(ws, rr, i, j):
    put(ws, f"A{rr}", f'="K{i}"&IF({SET}!B{FIN_ROW(i)}="",""," · "&{SET}!B{FIN_ROW(i)})', F_BOLD)
    put(ws, f"B{rr}", f"C{j}")
    put(ws, f"C{rr}", f'=IF({RUB}!B{RUB_ROW(j)}="","",{RUB}!B{RUB_ROW(j)})')


# ------------------------------------------------------------------ 02 Scores
ws = sheet("02 Scores", "02 · Scores before calibration",
           "One row per finalist and criterion. Enter each panelist's score, evidence basis and note after all have submitted. Never edit a score once entered.",
           [12, 6, 30, 7] + [7] * N_PAN + [14] * N_PAN + [36] * N_PAN, landscape=True)
put(ws, "A5", "Scores entered", F_BOLD)
put(ws, "C5", f'=IF({SET}!C13*{SET}!C22=0,"Name the panel and finalists on 00 Setup",COUNT(F{FIRST}:J{LAST})&" of "&({SET}!C13*{SET}!C22*{RUB}!D18))')
put(ws, "E5", f'=IF(COUNT(F{FIRST}:J{LAST})=0,"",IF(COUNT(F{FIRST}:J{LAST})>={SET}!C13*{SET}!C22*{RUB}!D18,"All in: ready to reveal","Waiting for scores"))', F_BOLD, border=False)
labels = ["Finalist", "Ref", "Criterion", "Weight"]
labels += [f'P{p} score' for p in range(1, N_PAN + 1)]
labels += [f'P{p} evidence basis' for p in range(1, N_PAN + 1)]
labels += [f'P{p} evidence note' for p in range(1, N_PAN + 1)]
head(ws, 6, labels)
for rr, i, j in rows_kc():
    kc_labels(ws, rr, i, j)
    put(ws, f"D{rr}", f'=IF({RUB}!D{RUB_ROW(j)}="","",{RUB}!D{RUB_ROW(j)})', align=CENTER)
    for c in PANEL:
        put(ws, f"{c}{rr}", None, fill=FILL_INPUT, align=CENTER)
    for c in BASIS + NOTES:
        put(ws, f"{c}{rr}", None, fill=FILL_INPUT)
validate(ws, f"F{FIRST}:J{LAST}", type="whole", operator="between", formula1="0", formula2="5")
validate(ws, f"K{FIRST}:O{LAST}", type="list", formula1='"Documented,Interview-stated,Inferred,None"')
validate(ws, f"P{FIRST}:T{LAST}", type="textLength", operator="greaterThanOrEqual", formula1="40",
         errorTitle="Evidence note too short", error="Name the source and the fact, in at least 40 characters.")
ws.freeze_panes = f"D{FIRST}"

# ------------------------------------------------------------------ 03 Divergence
ws = sheet("03 Divergence", "03 · Where the panel diverged",
           "Calculated from sheet 02. Scores of 0 mean no evidence and are left out of the spread and the mean. No totals appear before calibration.",
           [12, 6, 30, 10, 10, 8, 8, 8, 10, 12, 16], landscape=True)
head(ws, 6, ["Finalist", "Ref", "Criterion", "Scores entered", "Scoring > 0", "Min", "Max", "Spread", "Mean", "Provisional", "Flag"])
for rr, i, j in rows_kc():
    kc_labels(ws, rr, i, j)
    s = f"{SCO}!F{rr}:J{rr}"
    put(ws, f"D{rr}", f"=COUNT({s})", align=CENTER)
    put(ws, f"E{rr}", f'=COUNTIF({s},">0")', align=CENTER)
    put(ws, f"F{rr}", f'=IF(E{rr}=0,"",SMALL({s},COUNTIF({s},0)+1))', align=CENTER)
    put(ws, f"G{rr}", f'=IF(E{rr}=0,"",MAX({s}))', align=CENTER)
    put(ws, f"H{rr}", f'=IF(E{rr}=0,"",G{rr}-F{rr})', align=CENTER)
    put(ws, f"I{rr}", f'=IF(E{rr}=0,"",AVERAGEIF({s},">0"))', align=CENTER, fmt="0.00")
    # METHOD §5: the panel mean, rounded to the nearest integer, ties rounding down.
    put(ws, f"J{rr}", f'=IF(E{rr}=0,"",ROUNDUP(I{rr}-0.5,0))', align=CENTER)
    put(ws, f"K{rr}", f'=IF(OR(C{rr}="",D{rr}=0),"",IF(E{rr}=0,"No evidence",IF(H{rr}>=3,"Contradiction",IF(H{rr}=2,"Divergence","Agreement"))))', F_BOLD)
put(ws, f"C{LAST + 2}", "Divergences (spread 2)", F_BOLD); put(ws, f"D{LAST + 2}", f'=COUNTIF(K{FIRST}:K{LAST},"Divergence")', align=CENTER)
put(ws, f"C{LAST + 3}", "Contradictions (spread 3 or 4)", F_BOLD); put(ws, f"D{LAST + 3}", f'=COUNTIF(K{FIRST}:K{LAST},"Contradiction")', align=CENTER)
put(ws, f"C{LAST + 4}", "No evidence", F_BOLD); put(ws, f"D{LAST + 4}", f'=COUNTIF(K{FIRST}:K{LAST},"No evidence")', align=CENTER)
put(ws, f"C{LAST + 5}", "Reviewer agreement (spread 0 or 1)", F_BOLD)
put(ws, f"D{LAST + 5}", f'=IF(COUNTIF(K{FIRST}:K{LAST},"?*")-D{LAST + 4}=0,"",COUNTIF(K{FIRST}:K{LAST},"Agreement")/(COUNTIF(K{FIRST}:K{LAST},"?*")-D{LAST + 4}))',
    align=CENTER, fmt="0%")
put(ws, f"E{LAST + 5}", f'=IF(D{LAST + 5}="","",IF(D{LAST + 5}>=0.8,"High",IF(D{LAST + 5}>=0.6,"Medium","Low")))', F_BOLD, border=False)
ws.freeze_panes = f"D{FIRST}"

# ------------------------------------------------------------------ 04 Calibration
ws = sheet("04 Calibration", "04 · Calibration",
           "Resolve every flagged cell in the session: an agreed score and a written reason naming the evidence that settled it. Unresolved cells keep the panel mean.",
           [12, 6, 30, 14, 8, 11, 11, 48, 9, 14], landscape=True)
head(ws, 6, ["Finalist", "Ref", "Criterion", "Flag", "Spread", "Provisional", "Calibrated score", "Reason recorded", "Final", "Status"])
for rr, i, j in rows_kc():
    kc_labels(ws, rr, i, j)
    put(ws, f"D{rr}", f"={DIV}!K{rr}")
    put(ws, f"E{rr}", f"={DIV}!H{rr}", align=CENTER)
    put(ws, f"F{rr}", f"={DIV}!J{rr}", align=CENTER)
    put(ws, f"G{rr}", None, fill=FILL_INPUT, align=CENTER)
    put(ws, f"H{rr}", None, fill=FILL_INPUT)
    flagged = f'OR(D{rr}="Divergence",D{rr}="Contradiction")'
    put(ws, f"I{rr}", f'=IF(D{rr}="","",IF(D{rr}="No evidence",0,IF({flagged},IF(G{rr}="",F{rr},G{rr}),F{rr})))', F_BOLD, align=CENTER)
    put(ws, f"J{rr}", f'=IF(D{rr}="","",IF(D{rr}="No evidence","No evidence",IF({flagged},IF(G{rr}="","Unresolved",IF(H{rr}="","Reason missing","Calibrated")),"Agreement")))', F_BOLD)
validate(ws, f"G{FIRST}:G{LAST}", type="whole", operator="between", formula1="0", formula2="5")
for k, (label, formula) in enumerate([
    ("Flagged cells", f'=COUNTIF(D{FIRST}:D{LAST},"Divergence")+COUNTIF(D{FIRST}:D{LAST},"Contradiction")'),
    ("Calibrated", f'=COUNTIF(J{FIRST}:J{LAST},"Calibrated")'),
    ("Unresolved", f'=COUNTIF(J{FIRST}:J{LAST},"Unresolved")'),
    ("Reason missing", f'=COUNTIF(J{FIRST}:J{LAST},"Reason missing")'),
    ("Calibrated score entered on an unflagged cell", f'=SUMPRODUCT((G{FIRST}:G{LAST}<>"")*(J{FIRST}:J{LAST}<>"Calibrated")*(J{FIRST}:J{LAST}<>"Reason missing"))'),
], start=2):
    put(ws, f"C{LAST + k}", label, F_BOLD); put(ws, f"D{LAST + k}", formula, align=CENTER)
ws.freeze_panes = f"D{FIRST}"

# ------------------------------------------------------------------ 05 Results
ws = sheet("05 Results", "05 · Scores after calibration and the weighted result",
           "Entry order only. Never sort, colour or label finalists. The weighted result (0 to 100) is advisory: an input to the recommendation, not the recommendation.",
           [7, 32, 9] + [13] * N_FIN, landscape=True)
fin_cols = [col(4 + i) for i in range(N_FIN)]  # D..I
head(ws, 6, ["Ref", "Criterion", "Weight %"] + [f"K{i}" for i in range(1, N_FIN + 1)])
for i, c in enumerate(fin_cols, 1):
    ws[f"{c}6"].value = f'="K{i}"&IF({SET}!B{FIN_ROW(i)}="",""," · "&{SET}!B{FIN_ROW(i)})'
for j in range(1, N_CRIT + 1):
    rr = 7 + j  # rows 8..16
    put(ws, f"A{rr}", f"C{j}", F_BOLD)
    put(ws, f"B{rr}", f'=IF({RUB}!B{RUB_ROW(j)}="","",{RUB}!B{RUB_ROW(j)})')
    put(ws, f"C{rr}", f'=IF({RUB}!D{RUB_ROW(j)}="","",{RUB}!D{RUB_ROW(j)})', align=CENTER)
    for i, c in enumerate(fin_cols, 1):
        src = FIRST + (i - 1) * N_CRIT + (j - 1)
        put(ws, f"{c}{rr}", f"={CAL}!I{src}", align=CENTER)
put(ws, "B17", "Weighted result (advisory)", F_BOLD)
put(ws, "B18", "Criteria with no evidence", F_BOLD)
put(ws, "B19", "Unresolved cells", F_BOLD)
put(ws, "B20", "Note", F_BOLD)
for i, c in enumerate(fin_cols, 1):
    blk = f"{CAL}!J{FIRST + (i - 1) * N_CRIT}:J{FIRST + i * N_CRIT - 1}"
    # METHOD §6: result = sum(weight x calibrated score) / 5, from 0 to 100.
    put(ws, f"{c}17", f'=IF(COUNT({c}8:{c}16)=0,"",SUMPRODUCT($C$8:$C$16,{c}8:{c}16)/5)', F_BOLD, align=CENTER, fmt="0.0")
    put(ws, f"{c}18", f'=IF({c}17="","",COUNTIF({blk},"No evidence"))', align=CENTER)
    put(ws, f"{c}19", f'=IF({c}17="","",COUNTIF({blk},"Unresolved"))', align=CENTER)
    put(ws, f"{c}20", f'=IF({c}17="","",IF({c}18>0,"Incomplete",""))', align=CENTER)
put(ws, "A22", "Internal checks for the Verdict. Not for the client, and never shown as an order of finalists.", F_NOTE, border=False)
put(ws, "B23", "Highest weighted result", F_BOLD)
put(ws, "C23", f'=IF(COUNT(D17:I17)=0,"",INDEX($D$6:$I$6,MATCH(MAX(D17:I17),D17:I17,0)))', F_BOLD)
put(ws, "B24", "Its criteria scored 1 or 2", F_BOLD)
put(ws, "C24", f'=IF(C23="","",SUMPRODUCT((INDEX($D$8:$I$16,0,MATCH(MAX(D17:I17),D17:I17,0))>0)*(INDEX($D$8:$I$16,0,MATCH(MAX(D17:I17),D17:I17,0))<=2)*($C$8:$C$16<>"")))')
ws.merge_cells("D24:I24")
put(ws, "D24", "A risk for section 07 of the Verdict. If the recommendation names a different finalist, section 08 needs a written job-related reason.", F_NOTE, border=False)
ws.freeze_panes = "D8"

# ------------------------------------------------------------------ 06 Outcomes
ws = sheet("06 Outcomes", "06 · Outcome register",
           "Recorded against the role, by initials only. It is the only way Calibre can one day make a claim about itself.", [14, 14, 13, 26, 14, 13, 26, 16, 26, 30], landscape=True)
put(ws, "A5", "Candidate hired", F_BOLD); put(ws, "C5", None, fill=FILL_INPUT)
put(ws, "E5", "Start date", F_BOLD); put(ws, "F5", None, fill=FILL_INPUT, fmt="d mmm yyyy")
head(ws, 7, ["Checkpoint", "Due", "Still in the seat?", "If not, why", "Manager satisfaction (1 to 5)", "Early performance concerns?",
             "One line", "Did the Verdict flag what later mattered?", "One line", "Change to the rubric or method"])
for rr, days in zip(range(8, 11), [90, 180, 365]):
    put(ws, f"A{rr}", f"{days} days", F_BOLD)
    put(ws, f"B{rr}", f'=IF($F$5="","",$F$5+{days})', fmt="d mmm yyyy")
    for c in "CDEFGHIJ":
        put(ws, f"{c}{rr}", None, fill=FILL_INPUT, align=CENTER if c in "CEFH" else WRAP)
validate(ws, "C8:C10", type="list", formula1='"Yes,No"')
validate(ws, "F8:F10", type="list", formula1='"Yes,No"')
validate(ws, "H8:H10", type="list", formula1='"Yes,Partly,No"')
validate(ws, "E8:E10", type="whole", operator="between", formula1="1", formula2="5")

def fit_rows(ws):
    """Set row heights for wrapped text; neither Excel nor LibreOffice sizes rows written by openpyxl."""
    import math
    spans = {}
    for rng in ws.merged_cells.ranges:
        width = sum(ws.column_dimensions[col(c)].width or 9 for c in range(rng.min_col, rng.max_col + 1))
        spans[(rng.min_row, rng.min_col)] = width
    for row in ws.iter_rows(min_row=4):
        lines = 1
        for c in row:
            if isinstance(c.value, str) and not c.value.startswith("="):
                width = spans.get((c.row, c.column), ws.column_dimensions[col(c.column)].width or 9)
                size = c.font.size or 10
                chars = max(width * 1.2 * 10 / size, 1)
                lines = max(lines, sum(math.ceil(max(len(part), 1) / chars) for part in c.value.split("\n")))
        ws.row_dimensions[row[0].row].height = max(15, lines * 13.5)


for w in wb.worksheets:
    w.sheet_properties.tabColor = FOREST600
    fit_rows(w)
wb.save(OUT)
print("wrote", OUT)
