"""Builds kit/Calibre_Engagement_Time_Log.xlsx in the Calibre brand. Internal: never sent to a client.

One time log for the one product (DEC-2026-02). Budgets are the old panel-scope budgets, with the steps
renamed to product/METHOD.md. The economics test the fee against the founder-hour hurdle, all-in.

Usage: python3 kit/build/build_time_log.py [out.xlsx]
Then recalculate (LibreOffice) so cached values exist, e.g. the xlsx skill's scripts/recalc.py.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE.parent / "Calibre_Engagement_Time_Log.xlsx"

# Calibre tokens, brand/HANDBOOK.html §03. Same as build_working_file.py.
FOREST50, FOREST100, FOREST600, FOREST700, FOREST900 = "F0F5F3", "D5E3DE", "265147", "1C3D35", "0E211B"
STONE600, STONE500, WHITE = "394740", "596B66", "FFFFFF"
SERIF, SANS = "Lora", "Plus Jakarta Sans"
USE_STATEMENT = ("Structured hiring advisory for finance roles. Not a psychometric assessment. Does not predict "
                 "performance. The decision, and its consequences, rest with the employer.")

thin = Side(style="thin", color=FOREST100)
rule = Side(style="medium", color=FOREST600)
F_BODY = Font(name=SANS, size=10, color=STONE600)
F_BOLD = Font(name=SANS, size=10, bold=True, color=FOREST700)
F_HEAD = Font(name=SANS, size=9, bold=True, color=WHITE)
F_NOTE = Font(name=SANS, size=9, italic=True, color=STONE500)
F_TITLE = Font(name=SERIF, size=18, color=FOREST900)
F_SUB = Font(name=SERIF, size=13, color=FOREST900)
F_EYEBROW = Font(name=SANS, size=8, bold=True, color=FOREST600)
FILL_HEAD = PatternFill("solid", fgColor=FOREST900)
FILL_INPUT = PatternFill("solid", fgColor=FOREST50)
FILL_TOTAL = PatternFill("solid", fgColor=FOREST100)
WRAP = Alignment(wrap_text=True, vertical="top")
RIGHT = Alignment(horizontal="right", vertical="top")
HRS, USD, VAR = "0.00", "#,##0", "+0.00;-0.00;0.00"

wb = Workbook()
wb.remove(wb.active)


def sheet(name, heading, intro, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws["A1"], ws["A1"].font = "CALIBRE BY DAFTAR  ·  INTERNAL", F_EYEBROW
    ws["A2"], ws["A2"].font = heading, F_TITLE
    ws["A3"], ws["A3"].font = intro, F_NOTE
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[col(i)].width = w
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def head(ws, row, labels, start=2):
    for i, text in enumerate(labels):
        c = ws.cell(row, start + i, text)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, WRAP, Border(bottom=rule)


def put(ws, ref, value, font=F_BODY, fill=None, align=WRAP, fmt=None):
    c = ws[ref]
    c.value, c.font, c.alignment, c.border = value, font, align, Border(bottom=thin)
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def hours_input(ws, rng):
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="40", allow_blank=True,
                        showErrorMessage=True, errorTitle="Hours", error="Enter hours as a number from 0 to 40.")
    ws.add_data_validation(dv)
    for part in rng.split():
        dv.add(part)


# ---------------------------------------------------------------- Read me
ws = sheet("Read me", "Engagement time log",
           "Internal. Never sent to a client. One copy per engagement.", [3, 100])
lines = [
    ("What this is", F_SUB),
    ("The budget is the scope design. Log all-in founder hours against it, including selling and admin, "
     "and test the fee against the founder-hour hurdle. A Verdict that does not clear the hurdle at actual "
     "is a scope problem to fix before the next sale, not a rounding error.", F_BODY),
    ("How to use it", F_SUB),
    ("1. Assumptions: confirm the fee and the hurdle. Set a management hour cap only once the founder has decided one.", F_BODY),
    ("2. 01 Time log: enter actual hours as each step closes. Shaded cells are inputs; everything else is a formula.", F_BODY),
    ("3. Write a note against every step that overruns, naming the cause.", F_BODY),
    ("4. 02 Register: when the engagement closes, copy its fee and all-in actual hours into the register.", F_BODY),
    ("Scope", F_SUB),
    ("The fee covers four finalists. A fifth or sixth finalist, a second role or a rubric change after lock is "
     "quoted as a written variation; log that time on the variation's own copy of this file, not here.", F_BODY),
    ("The budgets are carried from the earlier panel-calibration time log. The live readout had no budget "
     "there, so it is budgeted at zero here: log its actual hours.", F_BODY),
    ("Use statement", F_SUB),
    (USE_STATEMENT, F_BODY),
]
for i, (text, font) in enumerate(lines, 5):
    ws.cell(i, 2, text).font = font
    ws.cell(i, 2).alignment = WRAP
    ws.row_dimensions[i].height = 22 if font is F_SUB else max(15, 15 * (len(text) // 110 + 1))

# ---------------------------------------------------------------- Assumptions
ws = sheet("Assumptions", "Assumptions",
           "The hurdle is the all-in opportunity cost of one founder hour. Shaded cells are inputs.", [3, 40, 14, 62])
head(ws, 5, ["Input", "Value", "Basis"])
inputs = [
    ("Founder hour hurdle (USD)", 350, USD, "All-in opportunity cost per founder hour (D/01 benchmark)."),
    ("Calibre Verdict fee (USD)", 6000, USD, "Rate card, founder decision 5 Oct 2026. Rationale in product/PRICING.md."),
    ("Management hour cap (all-in)", None, HRS, "Not set. A founder decision; leave blank until made."),
]
for r, (label, value, fmt, basis) in enumerate(inputs, 6):
    put(ws, f"B{r}", label, F_BOLD)
    put(ws, f"C{r}", value, fill=FILL_INPUT, fmt=fmt, align=RIGHT)
    put(ws, f"D{r}", basis, F_NOTE)
hours_input(ws, "C8")
ws["B10"], ws["B10"].font = "Derived", F_SUB
head(ws, 11, ["Measure", "Value", "Formula"])
derived = [
    ("Hours available at the hurdle", "=IF(OR(N(C6)=0,N(C7)=0),\"\",C7/C6)", HRS, "Fee divided by the hurdle."),
    ("Headroom, ceiling less cap", "=IF(OR(C8=\"\",C12=\"\"),\"\",C12-C8)", VAR,
     "Blank until a cap is set. Negative means the cap is above the hurdle ceiling."),
    ("Budgeted all-in hours", "='01 Time log'!C33", HRS, "From sheet 01."),
    ("Headroom at budget", "=IF(C12=\"\",\"\",C12-C14)", VAR,
     "Hours left before the hurdle is breached. Negative means the budget already fails it."),
]
for r, (label, f, fmt, note) in enumerate(derived, 12):
    put(ws, f"B{r}", label, F_BOLD)
    put(ws, f"C{r}", f, fmt=fmt, align=RIGHT)
    put(ws, f"D{r}", note, F_NOTE)

# ---------------------------------------------------------------- 01 Time log
ws = sheet("01 Time log", "Calibre Verdict time log",
           "Budget is the scope design. Log actuals as each step closes. Shaded cells are inputs.", [3, 52, 11, 11, 11, 46])
for r, label in enumerate(["Engagement reference", "Role", "Criteria locked on"], 4):
    put(ws, f"B{r}", label, F_BOLD)
    put(ws, f"C{r}", None, fill=FILL_INPUT)
    ws.merge_cells(f"C{r}:E{r}")
ws["C6"].number_format = "d mmm yyyy"
head(ws, 8, ["Step", "Budget hrs", "Actual hrs", "Variance", "Note"])

DELIVERY = [  # METHOD.md steps; budgets from the earlier panel-calibration time log
    ("Read · intake call and preparation", 1.5, ""),
    ("Read · rubric drafting and criteria lock", 1, ""),
    ("Read · structured interview guide", 1.5, ""),
    ("Read · evidence review, K1", 0.75, ""),
    ("Read · evidence review, K2", 0.75, ""),
    ("Read · evidence review, K3", 0.75, ""),
    ("Read · evidence review, K4", 0.75, ""),
    ("Score · panel score and evidence-note review", 1.5, ""),
    ("Compare · divergence review", 1.5, ""),
    ("Calibrate · session preparation", 1, ""),
    ("Calibrate · session (75 minutes)", 1.25, ""),
    ("Calibrate · calibrated results", 1, ""),
    ("Verdict · memo drafting", 2, ""),
    ("Verdict · live readout", 0, "Not budgeted in the earlier log. Log the actual."),
    ("Outcomes · 90, 180 and 365-day follow-up", 0.5, ""),
]
SELLING = [
    "Prospecting and warm outreach attributable to this engagement",
    "Proposal, rate card and scoping call",
    "Engagement letter and admin",
    "Invoicing and collection",
]


def step_row(r, label, budget, note=""):
    put(ws, f"B{r}", label)
    put(ws, f"C{r}", budget, fmt=HRS, align=RIGHT)
    put(ws, f"D{r}", None, fill=FILL_INPUT, fmt=HRS, align=RIGHT)
    put(ws, f"E{r}", f"=IF(D{r}=\"\",\"\",D{r}-C{r})", fmt=VAR, align=RIGHT)
    put(ws, f"F{r}", note or None, F_NOTE if note else F_BODY, fill=None if note else FILL_INPUT)


def subtotal(r, label, first, last):
    put(ws, f"B{r}", label, F_BOLD, fill=FILL_TOTAL)
    put(ws, f"C{r}", f"=SUM(C{first}:C{last})", F_BOLD, FILL_TOTAL, RIGHT, HRS)
    put(ws, f"D{r}", f"=IF(COUNT(D{first}:D{last})=0,\"\",SUM(D{first}:D{last}))", F_BOLD, FILL_TOTAL, RIGHT, HRS)
    put(ws, f"E{r}", f"=IF(D{r}=\"\",\"\",D{r}-C{r})", F_BOLD, FILL_TOTAL, RIGHT, VAR)
    put(ws, f"F{r}", None, fill=FILL_TOTAL)


put(ws, "B9", "DELIVERY", F_EYEBROW)
D0 = 10
for i, (label, b, note) in enumerate(DELIVERY):
    step_row(D0 + i, label, b, note)
D1 = D0 + len(DELIVERY) - 1                                   # 24
subtotal(D1 + 1, "Delivery subtotal", D0, D1)                 # 25
put(ws, "B27", "SELLING AND ADMIN · UNBILLED", F_EYEBROW)
put(ws, "B28", "Budgeted at zero. Log the actuals: the hurdle is tested all-in, not on delivery alone.", F_NOTE)
ws.merge_cells("B28:F28")
S0 = 29
for i, label in enumerate(SELLING):
    step_row(S0 + i, label, 0)
S1 = S0 + len(SELLING) - 1                                     # 32
# All-in total sits at row 33 (Assumptions!C14 reads it).
put(ws, "B33", "ALL-IN TOTAL", F_BOLD, FILL_TOTAL)
put(ws, "C33", f"=C{D1 + 1}+SUM(C{S0}:C{S1})", F_BOLD, FILL_TOTAL, RIGHT, HRS)
put(ws, "D33", f"=IF(COUNT(D{D0}:D{D1},D{S0}:D{S1})=0,\"\",SUM(D{D0}:D{D1},D{S0}:D{S1}))", F_BOLD, FILL_TOTAL, RIGHT, HRS)
put(ws, "E33", "=IF(D33=\"\",\"\",D33-C33)", F_BOLD, FILL_TOTAL, RIGHT, VAR)
put(ws, "F33", None, fill=FILL_TOTAL)
hours_input(ws, f"D{D0}:D{D1} D{S0}:D{S1}")

ws["B35"], ws["B35"].font = "Economics", F_SUB
head(ws, 36, ["Measure", "At budget", "At actual", "", "Test"])
FEE, HURDLE, CAP, CEIL = "Assumptions!$C$7", "Assumptions!$C$6", "Assumptions!$C$8", "Assumptions!$C$12"
econ = [
    ("All-in founder hours", "=C33", "=D33", HRS, ""),
    ("Effective rate per hour (USD)", f"=IF(N(C37)=0,\"\",{FEE}/C37)", f"=IF(N(D37)=0,\"\",{FEE}/D37)", USD,
     "Must reach the hurdle on Assumptions."),
    ("Hours available at the hurdle", f"={CEIL}", f"=IF(D37=\"\",\"\",{CEIL})", HRS, "Fee divided by the hurdle."),
    ("Headroom to the hurdle (hours)", "=IF(C39=\"\",\"\",C39-C37)", "=IF(OR(D37=\"\",D39=\"\"),\"\",D39-D37)", VAR,
     "Negative means the hours exceed what the fee pays for."),
    ("Hours against management cap", f"=IF({CAP}=\"\",\"\",C37-{CAP})", f"=IF(OR({CAP}=\"\",D37=\"\"),\"\",D37-{CAP})",
     VAR, "Blank until a cap is set. Negative is under the cap."),
    ("Clears the hurdle?", f"=IF(C38=\"\",\"\",IF(C38>={HURDLE},\"Yes\",\"No\"))",
     f"=IF(D38=\"\",\"\",IF(D38>={HURDLE},\"Yes\",\"No\"))", None, "If No, tighten scope or reprice before the next sale."),
]
for r, (label, fb, fa, fmt, test) in enumerate(econ, 37):
    put(ws, f"B{r}", label, F_BOLD)
    put(ws, f"C{r}", fb, fmt=fmt, align=RIGHT)
    put(ws, f"D{r}", fa, fmt=fmt, align=RIGHT)
    put(ws, f"F{r}", test or None, F_NOTE)
ws.freeze_panes = "C9"

# ---------------------------------------------------------------- 02 Register
ws = sheet("02 Register", "Dry run and engagement register",
           "One row per dry run or paid engagement. After three paid Verdicts, review the fee and the budget against these actuals.",
           [3, 13, 18, 26, 11, 11, 11, 10, 44])
head(ws, 5, ["Date", "Engagement ref", "Role", "Fee (USD)", "All-in hrs", "Rate / hr", "Clears hurdle", "What overran, and the scope change made"])
for r in range(6, 21):
    put(ws, f"B{r}", None, fill=FILL_INPUT, fmt="d mmm yyyy")
    for c in "CD":
        put(ws, f"{c}{r}", None, fill=FILL_INPUT)
    put(ws, f"E{r}", None, fill=FILL_INPUT, fmt=USD, align=RIGHT)
    put(ws, f"F{r}", None, fill=FILL_INPUT, fmt=HRS, align=RIGHT)
    put(ws, f"G{r}", f"=IF(OR(E{r}=\"\",N(F{r})=0),\"\",E{r}/F{r})", fmt=USD, align=RIGHT)
    put(ws, f"H{r}", f"=IF(G{r}=\"\",\"\",IF(G{r}>={HURDLE},\"Yes\",\"No\"))", align=RIGHT)
    put(ws, f"I{r}", None, fill=FILL_INPUT)
hours_input(ws, "F6:F20")
r = 22
put(ws, f"B{r}", "Paid to date", F_BOLD, FILL_TOTAL)
ws.merge_cells(f"B{r}:D{r}")
put(ws, f"E{r}", "=IF(COUNT(E6:E20)=0,\"\",SUM(E6:E20))", F_BOLD, FILL_TOTAL, RIGHT, USD)
put(ws, f"F{r}", "=IF(COUNT(F6:F20)=0,\"\",SUM(F6:F20))", F_BOLD, FILL_TOTAL, RIGHT, HRS)
put(ws, f"G{r}", "=IF(OR(E22=\"\",N(F22)=0),\"\",E22/F22)", F_BOLD, FILL_TOTAL, RIGHT, USD)
put(ws, f"H{r}", f"=IF(G22=\"\",\"\",IF(G22>={HURDLE},\"Yes\",\"No\"))", F_BOLD, FILL_TOTAL, RIGHT)
put(ws, f"I{r}", "Blended rate across every logged engagement.", F_NOTE, FILL_TOTAL)
ws.freeze_panes = "B6"

wb.active = 1
wb.save(OUT)
print(f"wrote {OUT}")
